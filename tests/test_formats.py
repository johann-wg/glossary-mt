"""文書形式のアダプターのテスト."""

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from glossary_mt.evaluate import find_pairs
from glossary_mt.formats import open_document

MD = """\
---
title: 手順書
---
# 手順書

本文です。
- 箇条書き `code`
  1. 入れ子の番号付き
- [ ] チェック項目

```bash
echo "翻訳しない"
```

| 項目 | 内容 |
|---|---|
| 監視 | 閾値 a\\|b |
<!-- コメント -->
"""


def test_markdown_segments(tmp_path: Path):
    path = tmp_path / "doc.md"
    path.write_text(MD, encoding="utf-8")
    segs = {s.id: s for s in open_document(path).segments()}
    assert {sid: s.text for sid, s in segs.items()} == {
        "L4": "手順書",
        "L6": "本文です。",
        "L7": "箇条書き `code`",
        "L8": "入れ子の番号付き",
        "L9": "チェック項目",
        "L15c1": "項目",
        "L15c2": "内容",
        "L17c1": "監視",
        "L17c2": "閾値 a\\|b",
    }
    assert segs["L6"].context == "手順書"


def test_markdown_save_keeps_structure(tmp_path: Path):
    path = tmp_path / "doc.md"
    path.write_bytes(MD.replace("\n", "\r\n").encode("utf-8"))
    doc = open_document(path)
    out = tmp_path / "out.md"
    # 何も訳さなければバイト単位で同一 (改行コード CRLF も保持)
    doc.save({}, out)
    assert out.read_bytes() == path.read_bytes()

    doc.save({"L4": "절차서", "L8": "중첩된 번호", "L17c2": "임계치 a|b"}, out)
    lines = out.read_bytes().decode("utf-8").split("\r\n")
    assert lines[3] == "# 절차서"
    assert lines[7] == "  1. 중첩된 번호"
    assert lines[16] == "| 監視 | 임계치 a\\|b |"  # 表のセル内の | はエスケープする


def test_markdown_bom_is_preserved(tmp_path: Path):
    path = tmp_path / "bom.md"
    path.write_bytes(b"\xef\xbb\xbf# \xe8\xa6\x8b\xe5\x87\xba\xe3\x81\x97\n")
    doc = open_document(path)
    assert [s.text for s in doc.segments()] == ["見出し"]
    out = tmp_path / "out.md"
    doc.save({"L1": "제목"}, out)
    assert out.read_bytes() == "﻿# 제목\n".encode()


def test_excel_translates_only_text_cells(tmp_path: Path):
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append(["件名", "件数", "期限"])
    ws.append(["障害対応", 3, dt.date(2026, 11, 1)])
    ws["D1"] = "=COUNTA(A:A)"
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    wb.save(path)

    doc = open_document(path)
    segs = doc.segments()
    assert [(s.id, s.text) for s in segs] == [
        ("一覧!A1", "件名"),
        ("一覧!B1", "件数"),
        ("一覧!C1", "期限"),
        ("一覧!A2", "障害対応"),
    ]
    assert segs[3].context == "シート「一覧」 / 列「件名」"

    out = tmp_path / "out.xlsx"
    doc.save({"一覧!A1": "건명", "一覧!A2": "장애 대응"}, out)
    ws = openpyxl.load_workbook(out)["一覧"]
    assert ws["A1"].value == "건명" and ws["A1"].font.bold
    assert ws["A2"].value == "장애 대응"
    assert ws["B2"].value == 3
    assert ws["D1"].value == "=COUNTA(A:A)"


def test_powerpoint_roundtrip(tmp_path: Path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Pt

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "現状の課題"
    p = slide.placeholders[1].text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "障害の"
    run.font.bold = True
    run.font.size = Pt(20)
    p.add_run().text = "切り分け"
    slide.notes_slide.notes_text_frame.text = "ノートの文"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    doc = open_document(path)
    assert [(s.id, s.text) for s in doc.segments()] == [
        ("s1:1", "現状の課題"),
        ("s1:2", "障害の切り分け"),
        ("s1:3", "ノートの文"),
    ]
    out = tmp_path / "out.pptx"
    doc.save({"s1:2": "장애 원인 분석", "s1:3": "노트"}, out)
    slide = pptx.Presentation(str(out)).slides[0]
    p = slide.placeholders[1].text_frame.paragraphs[0]
    assert p.text == "장애 원인 분석"
    # 段落先頭の書式が訳文に引き継がれる
    assert p.runs[0].font.bold and p.runs[0].font.size == Pt(20)
    assert slide.notes_slide.notes_text_frame.text == "노트"


@pytest.mark.parametrize("split", ["dev", "test"])
def test_eval_data_references_have_same_structure(split):
    """参照訳は原文と同じ構造 (Segment の id が一致) で書かれている."""
    for src, ref, _ in find_pairs(f"data/{split}"):
        src_ids = [s.id for s in open_document(src).segments()]
        ref_ids = [s.id for s in open_document(ref).segments()]
        assert src_ids == ref_ids, src.name
