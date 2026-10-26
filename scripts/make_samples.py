"""評価データ・サンプルの Excel / PowerPoint を生成する.

バイナリのファイルを手で編集すると差分が追えないため、内容はこのスクリプトで管理する。
原文と参照訳は同じ関数で作るので、セルの位置 (= Segment の id) が必ず一致する。

    python scripts/make_samples.py
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def _sheet(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(row)
    for col, width in zip("ABCDEFGHIJ", widths, strict=False):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


# --- 課題管理表 (dev) ---------------------------------------------------------

ISSUES = [
    # No, 件名, 内容, 担当者, 期限, ステータス, 備考  (日本語, 韓国語)
    (
        1,
        ("出荷指示データの連携遅延", "출고 지시 데이터 연계 지연"),
        (
            "夜間バッチ終了後、配送センターへの出荷指示データの連携が最大 40 分遅れる。",
            "야간 배치 종료 후 물류센터로의 출고 지시 데이터 연계가 최대 40분 지연된다.",
        ),
        ("高橋", "다카하시"),
        dt.date(2026, 11, 5),
        ("対応中", "진행 중"),
        (
            "暫定対応として連携ジョブの開始時刻を 15 分前倒しした。",
            "임시 조치로 연계 잡의 시작 시각을 15분 앞당겼다.",
        ),
    ),
    (
        2,
        ("帳票の文字化け", "리포트 글자 깨짐"),
        (
            "一部の帳票で韓国語の住所が文字化けする。",
            "일부 리포트에서 한국어 주소의 글자가 깨진다.",
        ),
        ("김민수", "김민수"),
        dt.date(2026, 11, 10),
        ("対応中", "진행 중"),
        ("フォント設定の影響範囲を調査中。", "폰트 설정의 영향 범위를 조사 중."),
    ),
    (
        3,
        ("監視閾値の見直し", "모니터링 임계치 재검토"),
        (
            "CPU 使用率の監視閾値（80%）で誤検知が多い。",
            "CPU 사용률 모니터링 임계치(80%)에서 오탐이 많다.",
        ),
        ("佐藤", "사토"),
        dt.date(2026, 11, 12),
        ("未着手", "미착수"),
        (
            "お客様と閾値を 90% に変更する方向で調整する。",
            "고객사와 임계치를 90%로 변경하는 방향으로 조율한다.",
        ),
    ),
    (
        4,
        ("DB サーバーの冗長化", "DB 서버 이중화"),
        (
            "本番環境の DB サーバーが 1 台構成のため、障害時にサービスが停止する。",
            "운영 환경의 DB 서버가 1대 구성이기 때문에 장애 시 서비스가 중지된다.",
        ),
        ("이지은", "이지은"),
        dt.date(2026, 12, 20),
        ("未着手", "미착수"),
        ("恒久対応として冗長化構成を提案する。", "근본 조치로 이중화 구성을 제안한다."),
    ),
    (
        5,
        ("ジョブネットの再実行手順", "잡넷 재실행 절차"),
        (
            "夜間バッチ異常終了時の再実行手順が手順書に記載されていない。",
            "야간 배치 비정상 종료 시의 재실행 절차가 절차서에 기재되어 있지 않다.",
        ),
        ("佐藤", "사토"),
        dt.date(2026, 11, 7),
        ("完了", "완료"),
        ("手順書 v1.3 に追記済み。", "절차서 v1.3에 추가 완료."),
    ),
    (
        6,
        ("検証環境の証明書期限切れ", "스테이징 환경 인증서 만료"),
        (
            "検証環境の SSL 証明書が 11 月 30 日に期限切れになる。",
            "스테이징 환경의 SSL 인증서가 11월 30일에 만료된다.",
        ),
        ("高橋", "다카하시"),
        dt.date(2026, 11, 25),
        ("完了", "완료"),
        ("更新済み。（申し送り事項）", "갱신 완료.(인수인계 사항)"),
    ),
]
ISSUE_HEADERS = (
    ["No", "件名", "内容", "担当者", "期限", "ステータス", "備考"],
    ["No", "건명", "내용", "담당자", "기한", "상태", "비고"],
)
STATUS_SUMMARY = (("完了件数", "完了"), ("완료 건수", "완료"))


def issue_tracker(lang: int) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "課題一覧"
    rows = []
    for no, title, body, owner, due, status, note in ISSUES:
        rows.append([no, title[lang], body[lang], owner[lang], due, status[lang], note[lang]])
    _sheet(ws, ISSUE_HEADERS[lang], rows, [6, 28, 50, 12, 12, 12, 40])
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "yyyy-mm-dd"
    # 数式のセル: 翻訳では触れない. 参照先のステータスが訳されるため、数式の条件も言語ごとに作る
    label, status = STATUS_SUMMARY[lang]
    ws["I1"] = label
    ws["I1"].font = Font(bold=True)
    ws["J1"] = f'=COUNTIF(F2:F{len(ISSUES) + 1},"{status}")'
    dv = DataValidation(
        type="list", formula1='"対応中,未着手,完了"' if lang == 0 else '"진행 중,미착수,완료"'
    )
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(ISSUES) + 1}")
    return wb


# --- 監視設定パラメータシート (test) -----------------------------------------

MONITORS = [
    (
        "MON-001",
        ("CPU 使用率", "CPU 사용률"),
        ("90%", "90%"),
        ("警告", "경고"),
        ("運用チーム", "운영팀"),
        (
            "5 分間の平均値が閾値を超えた場合に通知する。",
            "5분간의 평균값이 임계치를 초과한 경우에 통지한다.",
        ),
    ),
    (
        "MON-002",
        ("メモリ使用率", "메모리 사용률"),
        ("85%", "85%"),
        ("警告", "경고"),
        ("運用チーム", "운영팀"),
        ("スワップの発生も合わせて監視する。", "스왑 발생도 함께 모니터링한다."),
    ),
    (
        "MON-003",
        ("ディスク使用率", "디스크 사용률"),
        ("80%", "80%"),
        ("重大", "심각"),
        ("運用チーム・お客様", "운영팀·고객사"),
        ("ログ領域 /var/log を対象とする。", "로그 영역 /var/log를 대상으로 한다."),
    ),
    (
        "MON-004",
        ("夜間バッチ処理時間", "야간 배치 처리 시간"),
        ("90 分", "90분"),
        ("重大", "심각"),
        ("運用チーム", "운영팀"),
        (
            "ジョブネット DMS_NIGHTLY の開始から終了までの時間。",
            "잡넷 DMS_NIGHTLY의 시작부터 종료까지의 시간.",
        ),
    ),
    (
        "MON-005",
        ("出荷指示連携の遅延", "출고 지시 연계 지연"),
        ("30 分", "30분"),
        ("重大", "심각"),
        ("運用チーム・配送センター", "운영팀·물류센터"),
        ("遅延時は暫定対応として手動連携を行う。", "지연 시에는 임시 조치로 수동 연계를 실시한다."),
    ),
    (
        "MON-006",
        ("証明書の有効期限", "인증서 유효기간"),
        ("30 日前", "30일 전"),
        ("警告", "경고"),
        ("運用チーム", "운영팀"),
        (
            "検証環境と本番環境の両方を対象とする。",
            "스테이징 환경과 운영 환경 모두를 대상으로 한다.",
        ),
    ),
    (
        "MON-007",
        ("帳票サーバーの死活監視", "리포트 서버 가동 모니터링"),
        ("3 回連続失敗", "3회 연속 실패"),
        ("重大", "심각"),
        ("運用チーム", "운영팀"),
        (
            "冗長化された 2 台のうち 1 台でも停止したら通知する。",
            "이중화된 2대 중 1대라도 중지되면 통지한다.",
        ),
    ),
]
MONITOR_HEADERS = (
    ["項目ID", "監視項目", "閾値", "重要度", "通知先", "説明"],
    ["항목 ID", "모니터링 항목", "임계치", "중요도", "통지처", "설명"],
)
HISTORY = [
    ("1.0", "2026/09/01", ("初版作成", "초판 작성"), ("高橋", "다카하시")),
    (
        "1.1",
        "2026/10/15",
        ("閾値の見直し（CPU 80% → 90%）", "임계치 재검토(CPU 80% → 90%)"),
        ("佐藤", "사토"),
    ),
]
HISTORY_HEADERS = (["版", "日付", "内容", "作成者"], ["판", "날짜", "내용", "작성자"])


def parameter_sheet(lang: int) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "監視設定"
    rows = [[mid, *(col[lang] for col in cols)] for mid, *cols in MONITORS]
    _sheet(ws, MONITOR_HEADERS[lang], rows, [10, 24, 14, 8, 24, 50])
    hist = wb.create_sheet("改訂履歴")
    rows = [[ver, date, body[lang], author[lang]] for ver, date, body, author in HISTORY]
    _sheet(hist, HISTORY_HEADERS[lang], rows, [6, 12, 40, 12])
    return wb


# --- 提案資料 (examples) --------------------------------------------------------


def proposal_deck(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "DMS 監視基盤の改善提案"
    slide.placeholders[1].text = "2026年11月 定例会議資料"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "現状の課題"
    body = slide.placeholders[1].text_frame
    body.text = "夜間バッチの障害検知から対応開始まで平均 25 分かかっている"
    for text in [
        "監視閾値が固定値のため、月末の繁忙期に誤検知が多い",
        "障害時の一次切り分けが担当者の経験に依存している",
    ]:
        body.add_paragraph().text = text
    p = body.add_paragraph()
    run = p.add_run()
    run.text = "影響範囲の把握に時間がかかり、お客様への第一報が遅れる"
    run.font.bold = True
    slide.notes_slide.notes_text_frame.text = (
        "前回の定例で合意した課題の再確認。数字は 9 月の実績。"
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "改善スケジュール"
    table = slide.shapes.add_table(4, 3, Inches(0.5), Inches(1.8), Inches(9), Inches(2)).table
    cells = [
        ("フェーズ", "内容", "期間"),
        ("1", "監視閾値の見直しと検証環境での確認", "11月"),
        ("2", "一次切り分け手順書の整備", "12月"),
        ("3", "本番環境への適用と申し送りルールの改訂", "2027年1月"),
    ]
    for r, row in enumerate(cells):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
            for para in table.cell(r, c).text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(14)
    prs.save(str(path))


def main() -> None:
    issue_tracker(0).save(ROOT / "data/dev/02_issue_tracker.ja.xlsx")
    issue_tracker(1).save(ROOT / "data/dev/02_issue_tracker.ko.xlsx")
    parameter_sheet(0).save(ROOT / "data/test/05_parameter_sheet.ja.xlsx")
    parameter_sheet(1).save(ROOT / "data/test/05_parameter_sheet.ko.xlsx")
    (ROOT / "examples").mkdir(exist_ok=True)
    proposal_deck(ROOT / "examples/proposal.ja.pptx")
    print("generated")


if __name__ == "__main__":
    main()
