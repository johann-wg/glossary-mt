"""Excel (.xlsx).

文字列のセルだけを訳し、数値・日付・数式のセルには触れない。
セルの値だけを置き換えるため、書式・列幅・結合・入力規則はそのまま残る。
シート名は数式から参照されている可能性があるため翻訳しない。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from glossary_mt.formats.base import Segment


class ExcelDocument:
    def __init__(self, path: Path):
        self._wb = openpyxl.load_workbook(path)
        self._cells = {}
        self._segments: list[Segment] = []
        for ws in self._wb.worksheets:
            headers = {c.column: c.value for c in ws[1] if isinstance(c.value, str)}
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or not value.strip() or value.startswith("="):
                        continue
                    sid = f"{ws.title}!{cell.coordinate}"
                    context = f"シート「{ws.title}」"
                    if cell.row > 1 and cell.column in headers:
                        context += f" / 列「{headers[cell.column]}」"
                    self._cells[sid] = cell
                    self._segments.append(Segment(sid, value, context))

    def segments(self) -> list[Segment]:
        return list(self._segments)

    def save(self, translations: dict[str, str], path: Path) -> None:
        for sid, text in translations.items():
            self._cells[sid].value = text
        self._wb.save(path)
