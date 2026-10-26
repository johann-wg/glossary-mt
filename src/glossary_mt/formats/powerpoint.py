"""PowerPoint (.pptx).

段落単位で訳す (図形・表・グループ内の図形・ノート)。
段落内で書式が異なる部分 (一部だけ太字など) は、訳すと語順が変わるため対応が取れない。
訳文には段落の先頭の書式を適用する。
"""

from __future__ import annotations

import copy
from pathlib import Path

from glossary_mt.formats.base import Segment

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError as e:  # pragma: no cover
    raise ImportError(
        "PowerPoint には python-pptx が必要です: pip install 'glossary-mt[pptx]'"
    ) from e


def _text_frames(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _text_frames(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


class PowerPointDocument:
    def __init__(self, path: Path):
        self._prs = Presentation(str(path))
        self._paragraphs = {}
        self._segments: list[Segment] = []
        for n, slide in enumerate(self._prs.slides, start=1):
            frames = list(_text_frames(slide.shapes))
            if slide.has_notes_slide:
                frames.append(slide.notes_slide.notes_text_frame)
            title = slide.shapes.title.text if slide.shapes.title is not None else ""
            context = f"スライド {n}" + (f"「{title}」" if title else "")
            k = 0
            for frame in frames:
                for p in frame.paragraphs:
                    if not p.text.strip():
                        continue
                    k += 1
                    sid = f"s{n}:{k}"
                    self._paragraphs[sid] = p
                    self._segments.append(Segment(sid, p.text, context))

    def segments(self) -> list[Segment]:
        return list(self._segments)

    def save(self, translations: dict[str, str], path: Path) -> None:
        for sid, text in translations.items():
            p = self._paragraphs[sid]
            first = p.runs[0]._r.rPr if p.runs else None
            rpr = copy.deepcopy(first) if first is not None else None
            p.text = text  # 段落の書式 (pPr) は残り、run と改行だけが置き換わる
            if rpr is not None:
                for run in p.runs:
                    if run._r.rPr is not None:
                        run._r.remove(run._r.rPr)
                    run._r.insert(0, copy.deepcopy(rpr))
        self._prs.save(str(path))
